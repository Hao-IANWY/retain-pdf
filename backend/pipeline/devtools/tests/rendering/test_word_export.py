"""保留排版的 Word 导出。

这个模块此前**从「scripts/ → pipeline/」那次重命名起就跑不起来**：`exporter.py` 还在
import 已经改名的 `SCRIPTS_ROOT`，一 import 就 ImportError；`cli.py` 缺
`if __name__ == "__main__"`，`python -m` 跑完静悄悄什么都不做；声明在 pyproject 里的
`python-docx` 也没装进 venv。三件事指向同一个原因——**没有任何测试碰过它**，所以半年
没人知道它是死的。

所以这里第一优先钉的不是排版效果，而是「它还活着」：模块能 import、CLI 有入口、产物
真的写出来了、结构是我们要的那种（每页一节、背景图、绝对定位文本框、公式是 OMML 而
不是图片）。
"""

from __future__ import annotations

import importlib
import json
import pkgutil
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

PIPELINE_ROOT = Path(__file__).resolve().parents[3]

docx = pytest.importorskip("docx", reason="python-docx 未安装（它在 pyproject 里声明了）")
fitz = pytest.importorskip("fitz", reason="PyMuPDF 未安装")

from docx.oxml.ns import qn  # noqa: E402

MATH_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def test_every_module_imports():
    """整包 import 一遍。

    这一条就能挡住当初那次改名:`exporter` import 了一个已经不存在的名字,而它是
    整条链路的入口。
    """
    import retainpdf_pipeline.render.output.word as package

    failures = []
    for module in pkgutil.iter_modules(package.__path__):
        name = f"retainpdf_pipeline.render.output.word.{module.name}"
        try:
            importlib.import_module(name)
        except Exception as exc:  # noqa: BLE001 - 这里就是要把失败原因摊开
            failures.append(f"{name}: {type(exc).__name__}: {exc}")
    assert not failures, "有模块 import 不了:\n" + "\n".join(failures)


def test_no_module_still_imports_the_removed_python_docx_builder():
    """整个 pipeline 包里不该再有指向已删模块的 import。

    `math_omml` / `textboxes` / `document_builder` 这三个是 python-docx 时代的东西，
    文档生成搬去 retainpdf2doc 之后删掉了。删的时候漏了一个 `devtools/export_layout_docx.py`
    ——它 import 的是 `devtools.word_export.cli`，早就不存在，而没有任何测试碰过它，
    所以整个套件全绿也发现不了。
    """
    removed = ("math_omml", "textboxes", "document_builder", "devtools.word_export")
    offenders = []
    this_file = Path(__file__).resolve()
    for path in PIPELINE_ROOT.rglob("*.py"):
        # 跳过自己:这条测试的说明文字里就写着那几个模块名，会匹配到自身。
        if "__pycache__" in path.parts or path.resolve() == this_file:
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        for name in removed:
            # ocr/document_schema 下另有一个同名的 document_builder，那个还在用。
            if f"word.{name}" in text or f"word_export.{name}" in text or "devtools.word_export" in text:
                offenders.append(f"{path.relative_to(PIPELINE_ROOT)} → {name}")
    assert not offenders, "还有模块引用已删掉的 python-docx 实现：" + "; ".join(sorted(set(offenders)))


def test_cli_has_an_entry_point():
    """`python -m …cli` 必须真的执行 main()。

    少了这一段,命令跑完既没有产物也没有报错——排查时看起来像「成功但没输出」。
    """
    source = (PIPELINE_ROOT / "retainpdf_pipeline" / "render" / "output" / "word" / "cli.py").read_text()
    assert '__name__ == "__main__"' in source, "cli.py 没有入口,python -m 跑了等于没跑"


# ---------------------------------------------------------------------------
# 夹具：测试里现造的一页 job。
#
# 以前这里从本机 data/jobs 里挑「最新的」一个已翻译 job 裁成一页来用。代价是：CI 上
# 没有 data/jobs，整组永远 skip；本机则随你最近翻了哪本书而变——换一本首页是大标题
# 的书，断言的前提就不成立了（误报过一次）。
#
# 当初不肯合成数据的理由是「真实译文块有 80+ 个字段，手搓写不全」。但这里的断言只
# 经过两条路：`build_render_page_specs`（吃 bbox / 角色 / 原文译文这几个字段）和从
# 译文 PDF 读回字号行距（只认 PDF 里的文字 span）。所以现造的数据只要把每条断言要
# 分辨的那个性质**刻意**造出来即可，见 _BLOCKS / _RENDERED 旁的说明。
# ---------------------------------------------------------------------------
_PAGE_SIZE = (595.0, 842.0)

# (block_id, bbox, 是否标题, 原文, 译文)
#   - 一个标题块：font_roles 把标题判成 bold，凑出「字重混杂」（加粗 1/4）。
#   - 两个正文段：leading_em ≈ 0.91，`1 + leading_em` 远大于实测比值 1.289。
#   - 一个塞不下的窄框：译文比原文长得多，阅读器那套收敛会把它缩小。
_BLOCKS = (
    ("p001-b001", (72.0, 60.0, 523.0, 96.0), True,
     "A Synthetic Study of Layout Preservation", "版式保留的合成研究"),
    ("p001-b002", (72.0, 120.0, 523.0, 260.0), False,
     "The first paragraph describes the method in enough words to wrap over several lines. " * 3,
     "第一段用足够多的文字描述方法，使其在页面上折成好几行。" * 6),
    ("p001-b003", (72.0, 280.0, 523.0, 420.0), False,
     "The second paragraph reports results and is long enough to occupy multiple lines. " * 3,
     "第二段报告结果，同样需要足够长，才能在这里占据多行文字。" * 6),
    ("p001-b004", (72.0, 440.0, 300.0, 480.0), False,
     "A cramped box whose translation is far too long for the space of the source text.",
     "这是一个很挤的框，译文比原文占据的空间长得多，排版时必须缩小字号才能装下全部内容；"
     "原文只占了两行，译文却要排成三行以上才放得下。"),
)

# 「流水线渲染出来的译文 PDF」里每块真正排出来的 (字号, 基线间距)。
# 每块都比 spec 的上界小一圈、行距也和 `字号 × (1 + leading_em)` 对不上——两条读回
# 测试分辨的就是这两件事。标题只有一行，量不出行距（读回得到 0）。
_RENDERED = {
    "p001-b001": (20.0, 0.0),
    "p001-b002": (9.0, 12.0),
    "p001-b003": (9.5, 12.5),
    "p001-b004": (6.0, 7.5),
}


def _translated_item(block_id: str, bbox, is_title: bool, source: str, translated: str) -> dict:
    role = "title" if is_title else "body"
    return {
        "item_id": block_id,
        "page_idx": 0,
        "block_idx": int(block_id[-3:]) - 1,
        "block_kind": "text",
        "block_type": "title" if is_title else "text",
        "layout_role": "title" if is_title else "paragraph",
        "semantic_role": role,
        "structure_role": role,
        "policy_translate": True,
        "bbox": list(bbox),
        "source_text": source,
        "protected_source_text": source,
        "translated_text": translated,
        "protected_translated_text": translated,
    }


def _write_rendered_lines(page, bbox, text: str, size: float, step: float) -> None:
    """按给定字号和基线间距把译文逐行写进框里，模拟 Typst 收敛后的排版结果。"""
    x0, y0, x1, y1 = bbox
    per_line = max(1, int((x1 - x0) // size))  # 中文 1em 等宽
    lines = [text[i:i + per_line] for i in range(0, len(text), per_line)]
    baseline = y0 + size
    for line in lines:
        assert baseline <= y1, "夹具算错了：这块的字排出框外了，读回会把它当成别的块"
        page.insert_text((x0, baseline), line, fontsize=size, fontname="china-s")
        baseline += step


@pytest.fixture
def tiny_job(tmp_path: Path) -> Path:
    """一页 job：源 PDF + 译文 JSON/manifest + 流水线「渲染出来」的译文 PDF。"""
    job_root = tmp_path / "job"
    for name in ("source", "translated", "rendered"):
        (job_root / name).mkdir(parents=True)

    with fitz.open() as source:
        page = source.new_page(width=_PAGE_SIZE[0], height=_PAGE_SIZE[1])
        for _block_id, bbox, is_title, text, _translated in _BLOCKS:
            page.insert_textbox(fitz.Rect(*bbox), text, fontsize=16 if is_title else 10)
        source.save(job_root / "source" / "synthetic.pdf")

    items = [_translated_item(*block) for block in _BLOCKS]
    page_name = "page-001-synthetic.json"
    (job_root / "translated" / page_name).write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")
    (job_root / "translated" / "translation-manifest.json").write_text(
        json.dumps({"pages": [{"page_index": 0, "path": page_name}]}), encoding="utf-8",
    )

    with fitz.open() as rendered:
        page = rendered.new_page(width=_PAGE_SIZE[0], height=_PAGE_SIZE[1])
        for block_id, bbox, _is_title, _source, translated in _BLOCKS:
            size, step = _RENDERED[block_id]
            _write_rendered_lines(page, bbox, translated, size, step)
        rendered.save(job_root / "rendered" / "synthetic-translated.pdf")
    return job_root


@pytest.fixture
def docx_toolchain() -> None:
    """真导出 .docx 要 node + 构建好的 retainpdf2doc；缺了就跳过，并说清楚缺什么。

    CI 的 python-services 不构建 Node 包，所以真导出那几条在 CI 上照旧跳过；但只
    走到规格层（`build_layout_spec`）的用例不需要它，现在在哪儿都跑。
    """
    from retainpdf_pipeline.render.output.word import exporter

    try:
        exporter._resolve_node()
        exporter._resolve_cli()
    except exporter.LayoutDocxToolchainError as exc:
        pytest.skip(f"Word 导出工具链不可用：{exc}")


@pytest.fixture
def tiny_job_without_render(tiny_job: Path) -> Path:
    """同一份 job，但没有译文 PDF。

    读不回来的时候导出还得能跑完，并且退回 spec 的值——这条路在真实场景里很常见
    （只跑了翻译还没渲染，或者产物被清掉了）。
    """
    shutil.rmtree(tiny_job / "rendered", ignore_errors=True)
    return tiny_job


def _export(job_root: Path, out: Path, **kwargs):
    from retainpdf_pipeline.render.output.word.exporter import export_layout_docx

    return export_layout_docx(
        job_root=job_root, output_path=out, dpi=kwargs.pop("dpi", 72), **kwargs,
    )


@pytest.mark.usefixtures("docx_toolchain")
def test_export_writes_a_docx(tiny_job: Path, tmp_path: Path):
    out = tmp_path / "layout.docx"
    result = _export(tiny_job, out)
    assert Path(result).exists(), "导出声称成功却没有产物"
    assert out.stat().st_size > 0


@pytest.mark.usefixtures("docx_toolchain")
def test_each_page_becomes_its_own_section_at_the_source_size(tiny_job: Path, tmp_path: Path):
    """节 = 页，且页面尺寸取**源 PDF 的**，不是 Word 默认的 Letter。

    写死一个具体尺寸会把测试钉在这份夹具上；真正要保证的是「和源文档一致」。
    """
    source = sorted((tiny_job / "source").glob("*.pdf"))[0]
    with fitz.open(source) as doc:
        rect = doc[0].rect

    out = tmp_path / "layout.docx"
    _export(tiny_job, out)
    document = docx.Document(str(out))
    assert len(document.sections) == 1
    section = document.sections[0]
    assert round(section.page_width.pt) == round(rect.width)
    assert round(section.page_height.pt) == round(rect.height)


@pytest.mark.usefixtures("docx_toolchain")
def test_translated_text_lands_in_absolute_textboxes(tiny_job: Path, tmp_path: Path):
    """译文必须在文本框里。

    直接写成段落的话，Word 会按流式重排，「保留排版」就不成立了。
    """
    out = tmp_path / "layout.docx"
    _export(tiny_job, out)
    body = docx.Document(str(out)).element.body
    boxes = body.findall(".//" + qn("w:txbxContent"))
    assert len(boxes) >= 2, f"只有 {len(boxes)} 个文本框"

    texts = [t.text for t in body.iter() if t.tag == qn("w:t") and t.text and t.text.strip()]
    assert texts, "文本框里一个字都没有"


@pytest.mark.usefixtures("docx_toolchain")
def test_the_source_page_is_kept_as_a_background_image(tiny_job: Path, tmp_path: Path):
    """底图是「保留排版」的另一半:线条、图、表格都靠它。"""
    out = tmp_path / "layout.docx"
    _export(tiny_job, out)
    body = docx.Document(str(out)).element.body
    assert body.findall(".//" + qn("a:blip")), "没有背景图"


@pytest.mark.usefixtures("docx_toolchain")
def test_upstream_already_drops_blocks_without_text(tiny_job: Path, tmp_path: Path):
    """空译文的块进不到 exporter 这一层。

    起初我给 exporter 里那句 `if not block.plain_text.strip(): continue` 写了一条
    「空块不该变成空文本框」的测试，反证时发现改坏它测试照样全绿。查下来原因是
    **build_render_page_specs 已经把空块滤掉了**，那句 continue 从来没执行过——
    任何针对它的断言都恒真。

    所以这里钉的是真正成立的那件事:上游保证了这一点。它哪天不再保证，这条会红，
    到时候 exporter 那句防御才真的开始起作用。
    """
    from retainpdf_pipeline.render.output.word.job_io import single_pdf, translated_pages
    from retainpdf_pipeline.render.layout.page_specs import build_render_page_specs

    page_path = next((tiny_job / "translated").glob("page-*.json"))
    blocks = json.loads(page_path.read_text(encoding="utf-8"))

    blank = json.loads(json.dumps(blocks[0], ensure_ascii=False))
    blank["item_id"] = "p001-bblank"
    for key in list(blank):
        if key.endswith("translated_text"):
            blank[key] = ""
    blank["source_text"] = ""
    blocks.append(blank)
    page_path.write_text(json.dumps(blocks, ensure_ascii=False), encoding="utf-8")

    specs = build_render_page_specs(
        source_pdf_path=single_pdf(tiny_job / "source"),
        translated_pages=translated_pages(tiny_job),
    )
    # 上游自己就把没有文本的块滤掉了：喂进去 N+1 个，出来的没有一个是空的。
    assert all(b.plain_text.strip() for b in specs[0].blocks), "上游放了空块进来"
    expected = len(specs[0].blocks)

    out = tmp_path / "layout.docx"
    _export(tiny_job, out)
    body = docx.Document(str(out)).element.body
    assert len(body.findall(".//" + qn("w:txbxContent"))) == expected


@pytest.mark.usefixtures("docx_toolchain")
def test_max_pages_limits_the_export(tiny_job: Path, tmp_path: Path):
    out = tmp_path / "layout.docx"
    _export(tiny_job, out, max_pages=1)
    assert len(docx.Document(str(out)).sections) == 1


@pytest.mark.usefixtures("docx_toolchain")
def test_the_console_subcommand_runs_end_to_end(tiny_job: Path, tmp_path: Path):
    """走 `retainpdf-pipeline layout-docx` 这条真实路径。

    Rust 那边起的就是这个——`derived_artifacts` 用 `deps.pipeline_command` 拼命令行，
    和 `side-by-side-pdf` 同一套。直接调函数测不出「子命令没注册」「参数名对不上」
    这两类错，而它们恰好是接 API 时最容易犯的。
    """
    out = tmp_path / "console.docx"
    proc = subprocess.run(
        [sys.executable, "-m", "retainpdf_pipeline.entrypoints.console", "layout-docx",
         "--job-root", str(tiny_job), "--output-docx", str(out), "--dpi", "72"],
        cwd=str(PIPELINE_ROOT), capture_output=True, text=True, timeout=300,
    )
    assert proc.returncode == 0, proc.stderr[-800:]
    assert out.exists(), f"子命令没有产出文件\nstdout={proc.stdout}\nstderr={proc.stderr[-500:]}"


def test_the_subcommand_is_listed_in_usage():
    """没注册进 COMMANDS 的话，Rust 那边拿到的是 exit code 2 加一句 unknown command。"""
    from retainpdf_pipeline.entrypoints import console

    assert "layout-docx" in console.COMMANDS


def test_the_document_builder_is_wired_into_the_workspace(monkeypatch):
    """文档由 `backend/packages/retainpdf2doc`（Node）生成，它必须接进 npm workspace。

    这条钉的是「环境说得清楚」：以前 Python 侧自己用 python-docx 生成，装漏了依赖
    只会在运行时抛 ImportError；现在换成起子进程，没构建的话必须给出一句能照着做的
    错误，而不是让人对着非零退出码猜。

    以前这里还断言了 `_CLI_ENTRY.is_file()`，但它挂在本机真实 job 的夹具上，CI 上
    从没跑过；本机跑的时候它又和真导出那几条的前提是同一件事。现在「构建好没有」交给
    docx_toolchain 夹具判定（缺了就跳过并说明缺什么），这里只留不依赖构建产物的两条，
    在哪儿都跑。
    """
    import json

    from retainpdf_pipeline.render.output.word import exporter

    manifest = json.loads((PIPELINE_ROOT.parents[1] / "package.json").read_text(encoding="utf-8"))
    assert "backend/packages/retainpdf2doc" in manifest["workspaces"], (
        "retainpdf2doc 没注册进 npm workspaces，npm ci 不会装它的依赖"
    )

    # 没构建时的报错要自己说清楚怎么修。环境变量会绕过 _CLI_ENTRY，先清掉。
    monkeypatch.delenv(exporter._CLI_ENV_VAR, raising=False)
    monkeypatch.setattr(exporter, "_CLI_ENTRY", exporter._CLI_ENTRY.with_name("does-not-exist.mjs"))
    with pytest.raises(exporter.LayoutDocxToolchainError, match="npm run build"):
        exporter._resolve_cli()


def _observed(job_root: Path, block):
    """这个块在译文 PDF 里真正排出来的字号/行距。"""
    from retainpdf_pipeline.render.output.word.typography_readback import (
        converged_typography, open_translated_document, read_page_lines)

    document = open_translated_document(job_root)
    assert document is not None, "夹具里没带译文 PDF，读回路径测不到"
    try:
        lines = read_page_lines(document, 0)
        return converged_typography(lines, block.content_rect, len(block.plain_text.strip()))
    finally:
        document.close()


@pytest.mark.usefixtures("docx_toolchain")
def test_font_size_is_the_size_typst_converged_to_not_the_upper_bound(
    tiny_job: Path, tmp_path: Path,
):
    """字号要用译文 PDF 里真正排出来的那个，不是 spec 里的上界。

    `block.font_size_pt` 是交给 Typst `pdftr_fit_markdown` 的 `max_size`——Typst 会在
    `[fit_min_font_size_pt, font_size_pt]` 里二分找装得下的字号。照上界排版，凡是当初
    被缩过的块在 Word 里都会溢出。本仓真实 job 的前 5 页里有 10 个块被缩过，最狠的一个
    是 11.35pt → 7.84pt。
    """
    from retainpdf_pipeline.render.output.word.job_io import single_pdf, translated_pages
    from retainpdf_pipeline.render.layout.page_specs import build_render_page_specs

    specs = build_render_page_specs(
        source_pdf_path=single_pdf(tiny_job / "source"),
        translated_pages=translated_pages(tiny_job),
    )
    blocks = [b for b in specs[0].blocks if b.plain_text.strip()]
    shrunk = [
        (b, o) for b in blocks
        if (o := _observed(tiny_job, b)) is not None and b.font_size_pt - o.font_size_pt > 0.15
    ]
    assert shrunk, "夹具这一页没有任何块被缩过字号，这条分辨不出对错"

    out = tmp_path / "layout.docx"
    _export(tiny_job, out)
    body = docx.Document(str(out)).element.body
    sizes = {int(s.get(qn("w:val"))) for s in body.findall(".//" + qn("w:sz")) if s.get(qn("w:val"))}

    block, observed = max(shrunk, key=lambda pair: pair[0].font_size_pt - pair[1].font_size_pt)
    # w:sz 的单位是半磅，会取整；两个值取整后撞在一起就说明这块分辨不出来。
    converged = int(max(1.0, observed.font_size_pt) * 2)
    upper_bound = int(max(1.0, block.font_size_pt) * 2)
    assert converged != upper_bound, "这块缩得太少，半磅取整后两个值一样，换一块"
    assert converged in sizes, f"没有用收敛后的 {observed.font_size_pt}pt；出现的字号是 {sorted(sizes)}"
    assert upper_bound not in sizes, f"还在用上界 {block.font_size_pt}pt 排版，这块会溢出"


@pytest.mark.usefixtures("docx_toolchain")
def test_line_height_is_measured_from_the_rendered_baselines(tiny_job: Path, tmp_path: Path):
    """行距量自译文 PDF 相邻行的基线距离。

    此前是拿 spec 的 leading_em 折算成 `font_size_pt * (1 + leading_em)`。折算本身方向
    是对的（Typst 的 `par(leading:)` 是行间空隙，Word 的 `w:line` 是行高本身，不能照搬
    当倍数），但结果偏高很多——本仓真实 job 上按字符加权，折算值比真实行距平均高
    **2.37pt**，一行十二三磅就是高了两成，正文一长就顶出框外。
    """
    from retainpdf_pipeline.render.output.word.job_io import single_pdf, translated_pages
    from retainpdf_pipeline.render.layout.page_specs import build_render_page_specs

    specs = build_render_page_specs(
        source_pdf_path=single_pdf(tiny_job / "source"),
        translated_pages=translated_pages(tiny_job),
    )
    candidates = []
    for block in specs[0].blocks:
        if not block.plain_text.strip() or block.leading_em <= 0:
            continue
        observed = _observed(tiny_job, block)
        if observed is None or observed.line_step_pt <= 0:
            continue
        derived = int(max(1.0, block.font_size_pt * (1.0 + block.leading_em)) * 20)
        measured = int(max(1.0, observed.line_step_pt) * 20)
        if derived != measured:
            candidates.append((block, observed, derived, measured))
    assert candidates, "夹具这一页量不到和折算值不同的行距，这条分辨不出对错"

    out = tmp_path / "layout.docx"
    _export(tiny_job, out)
    body = docx.Document(str(out)).element.body
    values = {
        int(s.get(qn("w:line"))) for s in body.findall(".//" + qn("w:spacing")) if s.get(qn("w:line"))
    }

    _block, _observed_typography, derived, measured = max(
        candidates, key=lambda row: abs(row[2] - row[3]),
    )
    assert measured in values, f"没有用量到的行距 {measured / 20:.2f}pt；出现的是 {sorted(values)}"
    assert derived not in values, f"还在用折算的 {derived / 20:.2f}pt，正文会顶出框外"


def _spec_blocks(job_root: Path, **kwargs):
    from retainpdf_pipeline.render.output.word.exporter import build_layout_spec

    spec, _base = build_layout_spec(job_root=job_root, dpi=kwargs.pop("dpi", 72), **kwargs)
    return spec, [block for page in spec["pages"] for block in page["blocks"]]


def test_without_a_rendered_pdf_it_falls_back_to_the_html_fit_not_the_upper_bound(
    tiny_job_without_render: Path,
):
    """读不到译文 PDF 时走阅读器那套字号收敛，**不是**退回 spec 的上界。

    上界是 Typst 二分的起点而不是结果。跨 9 本书 295 个块实测，照上界排会有 44.6%
    的字符顶出框外；换成这套收敛之后降到 1.5%，而且误差方向是偏小。
    """
    from retainpdf_pipeline.render.output.word.html_fit import fitted_typography
    from retainpdf_pipeline.render.output.word.job_io import single_pdf, translated_pages
    from retainpdf_pipeline.render.layout.page_specs import build_render_page_specs

    layout = build_render_page_specs(
        source_pdf_path=single_pdf(tiny_job_without_render / "source"),
        translated_pages=translated_pages(tiny_job_without_render),
    )
    shrunk = {
        b.block_id: fitted_typography(b)[0] for b in layout[0].blocks
        if b.plain_text.strip() and b.font_size_pt - fitted_typography(b)[0] > 0.15
    }
    assert shrunk, "夹具这一页没有任何块被收敛算法缩小，这条分辨不出对错"
    upper = {b.block_id: b.font_size_pt for b in layout[0].blocks}

    _spec, blocks = _spec_blocks(tiny_job_without_render)
    checked = 0
    for block in blocks:
        if block["id"] not in shrunk:
            continue
        checked += 1
        assert block["fontSizePt"] == pytest.approx(shrunk[block["id"]], abs=0.01), block["id"]
        assert block["fontSizePt"] < upper[block["id"]] - 0.1, (
            f"{block['id']} 还在按上界 {upper[block['id']]}pt 排，这块会溢出"
        )
    assert checked, "规格里一个被缩过的块都没有"


def test_the_fallback_line_height_is_the_measured_ratio_not_one_plus_leading(
    tiny_job_without_render: Path,
):
    """兜底行距用实测比值 1.289，不是 `1 + leading_em`。

    折算方向看着合理（Typst 的 `par(leading:)` 是行间空隙，Word 的 `w:line` 是行高
    本身），但结果**系统性高 21%**:跨 9 本书 111 个块实测，真实行距是字号的 1.289 倍，
    而 `1 + leading_em` 给出 1.560。
    """
    from retainpdf_pipeline.render.output.word.html_fit import LINE_STEP_RATIO
    from retainpdf_pipeline.render.output.word.job_io import single_pdf, translated_pages
    from retainpdf_pipeline.render.layout.page_specs import build_render_page_specs

    layout = build_render_page_specs(
        source_pdf_path=single_pdf(tiny_job_without_render / "source"),
        translated_pages=translated_pages(tiny_job_without_render),
    )
    # 只看 `1 + leading_em` 和实测比值分得开的块。标题 / 作者行这类块排得紧，leading_em
    # 只有 0.28~0.29，`1 + leading_em` 本来就落在 1.289 附近甚至更小——在这些块上
    # 「比折算值小」不成立，也分辨不出用的是哪条公式（夹具换成首页是大标题的书就红了）。
    leading = {
        b.block_id: b.leading_em for b in layout[0].blocks
        if b.leading_em > 0 and 1.0 + b.leading_em > LINE_STEP_RATIO + 0.05
    }
    assert leading, "夹具这一页没有行距明显大于实测比值的块，这条分辨不出对错"

    _spec, blocks = _spec_blocks(tiny_job_without_render)
    checked = 0
    for block in blocks:
        if block["id"] not in leading:
            continue
        checked += 1
        assert block["lineStepPt"] == pytest.approx(block["fontSizePt"] * LINE_STEP_RATIO, abs=0.01)
        derived = block["fontSizePt"] * (1.0 + leading[block["id"]])
        assert block["lineStepPt"] < derived - 0.1, (
            f"{block['id']} 的行距还是按 (1+leading_em) 折算的，高约两成"
        )
    assert checked, "规格里一个带行距的块都没有"
    assert LINE_STEP_RATIO == 1.289, "改这个常数要重新跑实测，别凭感觉调"


def test_a_block_is_not_given_a_neighbours_font_size():
    """归属按「行的中心落在框内」，不是 PyMuPDF 的 clip。

    用 `clip=` 量的话，压在框边上的邻块文字会被一起裁进来——之前拿它量出「52.7% 的
    字号和 spec 对不上」，那个数字整个是假的。这条钉的是:一个框里读不到自己的字时，
    宁可返回 None 退回 spec，也不要拿邻居的字号顶上。
    """
    from retainpdf_pipeline.render.output.word.typography_readback import converged_typography, _ObservedLine

    neighbour = _ObservedLine(
        x0=0.0, y0=0.0, x1=100.0, y1=10.0, baseline=8.0, sizes=((20.0, 300),),
    )
    # 框在 (0,100)-(100,140)，邻居那一行的中心在 y=5，压根不在框里。
    assert converged_typography([neighbour], [0.0, 100.0, 100.0, 140.0], 200) is None

    # 就算落在框内，读到的字远少于这个块该有的量，也不该据此定字号。
    inside = _ObservedLine(
        x0=0.0, y0=100.0, x1=100.0, y1=110.0, baseline=108.0, sizes=((20.0, 5),),
    )
    assert converged_typography([inside], [0.0, 100.0, 100.0, 140.0], 200) is None
    assert converged_typography([inside], [0.0, 100.0, 100.0, 140.0], 8) is not None


@pytest.mark.usefixtures("docx_toolchain")
def test_bold_blocks_are_bold(tiny_job: Path, tmp_path: Path):
    """字重来自 font_weight。不接的话标题和正文一样粗，Word 里读不出层次。"""
    from retainpdf_pipeline.render.output.word.job_io import single_pdf, translated_pages
    from retainpdf_pipeline.render.layout.page_specs import build_render_page_specs

    specs = build_render_page_specs(
        source_pdf_path=single_pdf(tiny_job / "source"),
        translated_pages=translated_pages(tiny_job),
    )
    blocks = specs[0].blocks
    bold_blocks = [b for b in blocks if str(b.font_weight or "").lower() == "bold"]
    assert bold_blocks, "夹具里没有加粗块，测不到东西"

    out = tmp_path / "layout.docx"
    _export(tiny_job, out)
    body = docx.Document(str(out)).element.body
    bold_runs = body.findall(".//" + qn("w:b"))
    assert bold_runs, "一个加粗都没有"


@pytest.mark.usefixtures("docx_toolchain")
def test_regular_blocks_are_not_bold(tiny_job: Path, tmp_path: Path):
    """反过来也要成立——不能整篇都加粗。

    第一版比的是「加粗的 run 数 < 总 run 数」，反证时发现整篇强制加粗它照样绿:
    OMML 公式里的 run 走的是另一条路、永远不带 w:b，所以那个不等式恒成立。
    改成按**文本框**比:常规字重的块里不该出现加粗。
    """
    from retainpdf_pipeline.render.output.word.job_io import single_pdf, translated_pages
    from retainpdf_pipeline.render.layout.page_specs import build_render_page_specs

    specs = build_render_page_specs(
        source_pdf_path=single_pdf(tiny_job / "source"),
        translated_pages=translated_pages(tiny_job),
    )
    blocks = [b for b in specs[0].blocks if b.plain_text.strip()]
    expected_bold = sum(1 for b in blocks if str(b.font_weight or "").lower() == "bold")
    assert 0 < expected_bold < len(blocks), (
        f"夹具里字重不够混杂（{expected_bold}/{len(blocks)} 加粗），这条分辨不出对错"
    )

    out = tmp_path / "layout.docx"
    _export(tiny_job, out)
    body = docx.Document(str(out)).element.body
    boxes = body.findall(".//" + qn("w:txbxContent"))
    bold_boxes = sum(1 for box in boxes if box.findall(".//" + qn("w:b")))
    assert bold_boxes == expected_bold, (
        f"{bold_boxes} 个文本框加粗，排版层说的是 {expected_bold} 个"
    )


def test_first_line_indent_reaches_the_spec_even_though_this_fixture_has_none(tiny_job: Path):
    """首行缩进进了规格。

    这份夹具测不到它的**效果**——夹具里所有块的 first_line_indent_pt 都是 0，真实数据
    也常常如此（中文正文的两字缩进由翻译侧直接写进文本，不靠排版属性）。以前这条只能去 grep
    exporter.py 的源码，现在规格是个普通 dict，至少能钉住字段真的在、且值是照搬的。
    """
    from retainpdf_pipeline.render.output.word.job_io import single_pdf, translated_pages
    from retainpdf_pipeline.render.layout.page_specs import build_render_page_specs

    layout = build_render_page_specs(
        source_pdf_path=single_pdf(tiny_job / "source"), translated_pages=translated_pages(tiny_job),
    )
    expected = {b.block_id: round(b.first_line_indent_pt or 0.0, 3) for b in layout[0].blocks}
    _spec, blocks = _spec_blocks(tiny_job)
    assert blocks, "规格里没有块"
    for block in blocks:
        assert "firstLineIndentPt" in block, f"{block['id']} 的规格里没有首行缩进字段"
        assert block["firstLineIndentPt"] == expected[block["id"]]


@pytest.fixture
def translate_only_job(tiny_job: Path, tmp_path: Path) -> tuple[Path, Path]:
    """长得像 translate-only 任务的 job:自己的 source/ 是空的。

    这类任务复用上游 OCR 任务的产物，源 PDF 在**父任务**目录里（job 记录的
    `source_artifact_job_id` 指过去）。Rust 那边 `resolve_source_pdf` 从记录解析得到
    正确路径，而导出这边如果按 `job_root/source/` 自己找，就是 0 个 PDF。
    """
    upstream = tmp_path / "upstream-source"
    upstream.mkdir()
    moved = sorted((tiny_job / "source").glob("*.pdf"))[0]
    target = upstream / moved.name
    shutil.move(str(moved), target)
    assert not list((tiny_job / "source").glob("*.pdf")), "夹具的 source/ 没清空"
    return tiny_job, target


@pytest.mark.usefixtures("docx_toolchain")
def test_a_job_whose_source_pdf_lives_elsewhere_still_exports(
    translate_only_job: tuple[Path, Path], tmp_path: Path,
):
    """源 PDF 不在 job_root 下时，调用方给了路径就该用它。

    真实案例:20260918070141-be37aa（translate-only）。Rust 侧解析到了父任务里的源
    PDF、检查通过，Python 侧却按 job_root 重新找，找到 0 个直接抛异常——整条导出 500，
    而且子进程的 stderr 是丢弃的，报错里只有一句 "failed to build layout-docx"。
    """
    job_root, source_pdf = translate_only_job
    out = tmp_path / "translate-only.docx"
    _export(job_root, out, source_pdf=source_pdf)
    assert out.exists() and out.stat().st_size > 0
    assert docx.Document(str(out)).element.body.findall(".//" + qn("w:txbxContent")), "没有文本框"


@pytest.mark.usefixtures("docx_toolchain")
def test_without_the_resolved_path_such_a_job_fails_loudly(
    translate_only_job: tuple[Path, Path], tmp_path: Path,
):
    """不给路径时仍然按 job_root 找——找不到要明确报错，不是产出一份空文档。"""
    job_root, _source_pdf = translate_only_job
    with pytest.raises(RuntimeError, match="expected exactly one PDF"):
        _export(job_root, tmp_path / "boom.docx")


@pytest.mark.usefixtures("docx_toolchain")
def test_the_console_subcommand_forwards_the_resolved_paths(
    translate_only_job: tuple[Path, Path], tmp_path: Path,
):
    """Rust 起的是子命令，所以参数名对不上等于没修。"""
    job_root, source_pdf = translate_only_job
    out = tmp_path / "console-translate-only.docx"
    proc = subprocess.run(
        [sys.executable, "-m", "retainpdf_pipeline.entrypoints.console", "layout-docx",
         "--job-root", str(job_root), "--output-docx", str(out),
         "--source-pdf", str(source_pdf), "--dpi", "72"],
        cwd=str(PIPELINE_ROOT), capture_output=True, text=True, timeout=300,
    )
    assert proc.returncode == 0, proc.stderr[-800:]
    assert out.exists(), f"子命令没有产出文件\nstderr={proc.stderr[-500:]}"


def test_every_module_in_the_word_package_is_tracked_by_git():
    """这个包里的每个 .py 都必须在版本库里。

    `.gitignore` 有一条不带斜杠前缀的 `output/`，它匹配**任意深度**的同名目录——
    包括 `retainpdf_pipeline/render/output/` 这个真实源码包。解除忽略的 `!` 规则原来
    写的是改名前的 `services/pipeline/services/rendering/output/`，改名后没人更新。

    后果很隐蔽:已跟踪的文件照常工作，谁都看不出问题，但**新建的文件会被
    `git add -A` 静默跳过**，`git status` 里也不显示。v4.2.5 就是这么少推了
    `html_fit.py` 和 `cli.py`，本地全绿、CI 报 ModuleNotFoundError 才查出来。

    .gitignore 已经修好；这条测试盯着它别再退化。
    """
    import subprocess

    package = PIPELINE_ROOT / "retainpdf_pipeline" / "render" / "output" / "word"
    on_disk = {
        path.name for path in package.glob("*.py")
        if "__pycache__" not in path.parts
    }
    assert on_disk, "包目录是空的，这条测试分辨不出对错"

    result = subprocess.run(
        ["git", "ls-files", str(package.relative_to(PIPELINE_ROOT.parents[1]))],
        cwd=PIPELINE_ROOT.parents[1], capture_output=True, text=True, check=False,
    )
    if result.returncode != 0:
        pytest.skip("不在 git 检出里")
    tracked = {Path(line).name for line in result.stdout.split() if line.endswith(".py")}

    missing = sorted(on_disk - tracked)
    assert not missing, (
        f"这些模块没进版本库，CI 上会 ModuleNotFoundError：{missing}。"
        "多半又是 .gitignore 那条 `output/` 把它们吃了——检查解除忽略规则的路径。"
    )


def test_packaged_deployments_can_point_at_the_cli_and_node(tmp_path: Path, monkeypatch):
    """打包环境用环境变量告知 retainpdf2doc 和 node 的真实位置。

    `_CLI_ENTRY` 是按仓库布局往上数六层算出来的，这只在**检出**里成立。装进桌面应用
    或 Docker 之后这个包从 site-packages 跑，往上数会数到 python 运行时目录里，
    `dist/cli.mjs` 当然不存在——v4.2.5 的 Mac 应用就是这么导不出 Word 的（而且当时
    前端会在请求失败前就把文件创建出来，于是表现成"打开什么都没有的空白文档"）。
    """
    from retainpdf_pipeline.render.output.word import exporter

    cli = tmp_path / "cli.mjs"
    cli.write_text("// stub", encoding="utf-8")
    node = tmp_path / "node"
    node.write_text("#!/bin/sh\n", encoding="utf-8")

    monkeypatch.setenv(exporter._CLI_ENV_VAR, str(cli))
    monkeypatch.setenv(exporter._NODE_ENV_VAR, str(node))
    assert exporter._resolve_cli() == cli
    assert exporter._resolve_node() == str(node)

    # 指错了要明确报出来，不能悄悄退回仓库布局。
    monkeypatch.setenv(exporter._CLI_ENV_VAR, str(tmp_path / "nope.mjs"))
    with pytest.raises(exporter.LayoutDocxToolchainError, match="不存在"):
        exporter._resolve_cli()
    monkeypatch.setenv(exporter._NODE_ENV_VAR, str(tmp_path / "nope"))
    with pytest.raises(exporter.LayoutDocxToolchainError, match="不存在"):
        exporter._resolve_node()


def test_the_missing_toolchain_message_says_how_to_fix_it(monkeypatch):
    """两种环境各自的修法都要写在报错里。"""
    from retainpdf_pipeline.render.output.word import exporter

    monkeypatch.delenv(exporter._CLI_ENV_VAR, raising=False)
    monkeypatch.setattr(exporter, "_CLI_ENTRY", Path("/definitely/not/here/cli.mjs"))
    with pytest.raises(exporter.LayoutDocxToolchainError) as caught:
        exporter._resolve_cli()
    message = str(caught.value)
    assert "npm run build --workspace retainpdf2doc" in message, "没告诉仓库里怎么修"
    assert exporter._CLI_ENV_VAR in message, "没告诉打包环境怎么指定路径"
